"""ASCENT SAFETY isolated document extractor.

Run as:  python -I ascent_extract_worker.py <extension>
Reads raw bytes from stdin, writes exactly one JSON object to stdout.

This process never executes document content. It applies OS resource limits to
itself (address space, CPU seconds, no core dumps) before touching any parser;
the parent additionally enforces a wall-clock timeout and kills it on expiry.
It is process isolation only, NOT a sandbox: run it in a container/low-privilege
account with no network for stronger containment.
"""

from __future__ import annotations

import io
import json
import os
import re
import sys
import zipfile
from pathlib import PurePosixPath

MAX_CHARS = 250_000
MAX_PDF_PAGES = int(os.environ.get("ASCENT_WORKER_MAX_PDF_PAGES", "100"))
MAX_ZIP_ENTRIES = 2000
MAX_ZIP_UNCOMPRESSED = 40 * 1024 * 1024
MAX_ZIP_RATIO = 100
MAX_IMAGE_PIXELS = 25_000_000
OCR_TIMEOUT_SECONDS = int(os.environ.get("ASCENT_WORKER_OCR_TIMEOUT", "20"))

TEXT_EXTS = {".txt", ".eml", ".msg"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED_EXTS = TEXT_EXTS | IMAGE_EXTS | {".pdf", ".docx"}
_MAGIC = {
    ".pdf": (b"%PDF-",),
    ".docx": (b"PK\x03\x04",),
    ".png": (b"\x89PNG\r\n\x1a\n",),
    ".jpg": (b"\xff\xd8\xff",),
    ".jpeg": (b"\xff\xd8\xff",),
    ".bmp": (b"BM",),
    ".tif": (b"II*\x00", b"MM\x00*"),
    ".tiff": (b"II*\x00", b"MM\x00*"),
}
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_PDF_INDICATORS = [
    (re.compile(rb"/JavaScript\b|/JS\b"), "pdf_javascript"),
    (re.compile(rb"/Launch\b"), "pdf_launch_action"),
    (re.compile(rb"/OpenAction\b"), "pdf_open_action"),
    (re.compile(rb"/EmbeddedFile\b"), "pdf_embedded_file"),
]


def apply_limits() -> bool:
    """Lower this process's own limits. Returns False where `resource` is unavailable."""
    try:
        import resource
    except ImportError:
        return False
    mem = int(os.environ.get("ASCENT_WORKER_MEM_BYTES", str(1 << 30)))
    cpu = int(os.environ.get("ASCENT_WORKER_CPU_SECONDS", "20"))
    for limit, value in (
        (resource.RLIMIT_AS, mem),
        (resource.RLIMIT_CPU, cpu),
        (resource.RLIMIT_CORE, 0),
    ):
        try:
            resource.setrlimit(limit, (value, value))
        except (ValueError, OSError):
            pass
    return True


def validate_content(ext: str, raw: bytes) -> tuple[str, str] | None:
    """Cheap pre-parse checks. Returns (error_code, message) or None when acceptable."""
    if ext == ".webp":
        if not (raw[:4] == b"RIFF" and raw[8:12] == b"WEBP"):
            return "content_mismatch", "File content does not match the .webp extension."
    elif ext in _MAGIC:
        head = raw[:1024] if ext == ".pdf" else raw[:16]
        if not any((sig in head) if ext == ".pdf" else head.startswith(sig) for sig in _MAGIC[ext]):
            return "content_mismatch", f"File content does not match the {ext} extension."
    elif ext in TEXT_EXTS:
        if ext == ".msg" and raw.startswith(_OLE_MAGIC):
            return (
                "unsupported_binary_msg",
                "Binary Outlook .msg files are not supported; export the message as .eml.",
            )
        if b"\x00" in raw[:8192]:
            return "binary_text", "File contains binary data and is not valid text."
    else:
        return "unsupported_type", f"Unsupported file type {ext!r}."
    return None


def _result(
    status: str,
    text: str = "",
    notes: list[str] | None = None,
    indicators: list[str] | None = None,
    error_code: str | None = None,
) -> dict:
    return {
        "status": status,
        "text": text[:MAX_CHARS],
        "notes": notes or [],
        "indicators": sorted(set(indicators or [])),
        "error_code": error_code,
    }


def decode_text(raw: bytes) -> tuple[str, list[str]]:
    notes: list[str] = []
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return raw.decode("utf-16"), notes
        except UnicodeDecodeError:
            pass
    try:
        return raw.decode("utf-8-sig"), notes
    except UnicodeDecodeError:
        notes.append("Content is not valid UTF-8; undecodable bytes were replaced.")
        return raw.decode("utf-8", errors="replace"), notes


def _extract_eml(raw: bytes) -> dict:
    from email import policy
    from email.parser import BytesParser

    try:
        msg = BytesParser(policy=policy.default).parsebytes(raw)
        parts = [f"{k}: {msg.get(k)}" for k in ("subject", "from", "to", "reply-to") if msg.get(k)]
        body = msg.get_body(preferencelist=("plain",))
        if body:
            parts.append(str(body.get_content()))
        else:
            payload = msg.get_payload(decode=True)
            if payload:
                parts.append(payload.decode("utf-8", errors="replace"))
        names = [p.get_filename() for p in msg.iter_attachments() if p.get_filename()]
        if names:
            parts.append("Attachments: " + ", ".join(names[:50]))
        return _result("ok", "\n".join(parts))
    except Exception:
        text, notes = decode_text(raw)
        notes.append("Email parsing was incomplete; decoded raw text was inspected.")
        return _result("ok", text, notes)


def _extract_pdf(raw: bytes) -> dict:
    indicators = [name for rx, name in _PDF_INDICATORS if rx.search(raw)]
    notes: list[str] = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw), strict=False)
        if getattr(reader, "is_encrypted", False):
            try:
                if reader.decrypt("") == 0:
                    return _result(
                        "failed",
                        notes=["Encrypted PDF could not be inspected without a password."],
                        indicators=indicators,
                        error_code="encrypted",
                    )
            except Exception:
                return _result(
                    "failed",
                    notes=["Encrypted PDF could not be inspected."],
                    indicators=indicators,
                    error_code="encrypted",
                )
        if len(reader.pages) > MAX_PDF_PAGES:
            return _result(
                "limit",
                notes=[f"PDF has more than {MAX_PDF_PAGES} pages."],
                indicators=indicators,
                error_code="too_many_pages",
            )
        chunks, total = [], 0
        for page in reader.pages:
            piece = page.extract_text() or ""
            chunks.append(piece)
            total += len(piece)
            if total > MAX_CHARS:
                notes.append("Extracted text was truncated at the size limit.")
                break
        text = "\n".join(chunks)
    except Exception:
        return _result(
            "failed",
            notes=["PDF extraction failed; file may be damaged, encrypted, or unsupported."],
            indicators=indicators,
            error_code="parse_error",
        )
    if not text.strip():
        notes.append(
            "PDF contains no extractable text; it may be scanned (OCR is not applied to PDFs)."
        )
        return _result("no_text", text, notes, indicators)
    return _result("ok", text, notes, indicators)


def _extract_docx(raw: bytes) -> dict:
    indicators: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            infos = zf.infolist()
            if (
                len(infos) > MAX_ZIP_ENTRIES
                or sum(i.file_size for i in infos) > MAX_ZIP_UNCOMPRESSED
            ):
                return _result(
                    "limit",
                    notes=["DOCX archive expands beyond the inspection limit."],
                    error_code="archive_too_large",
                )
            for info in infos:
                parts = PurePosixPath(info.filename.replace("\\", "/")).parts
                if info.filename.startswith(("/", "\\")) or ".." in parts:
                    return _result(
                        "limit",
                        notes=["DOCX archive has unsafe entry paths."],
                        error_code="archive_bad_paths",
                    )
                if (
                    info.file_size > (1 << 20)
                    and info.file_size / max(info.compress_size, 1) > MAX_ZIP_RATIO
                ):
                    return _result(
                        "limit",
                        notes=["DOCX archive entry has a suspicious compression ratio."],
                        error_code="archive_too_large",
                    )
                lowered = info.filename.lower()
                if lowered.endswith("vbaproject.bin"):
                    indicators.append("docx_vba_project")
                elif "oleobject" in lowered or lowered.endswith("activex.xml"):
                    indicators.append("docx_embedded_object")
                elif lowered.endswith(".rels") and info.file_size <= (1 << 20):
                    if b'TargetMode="External"' in zf.read(info):
                        indicators.append("docx_external_relationship")
    except zipfile.BadZipFile:
        return _result("failed", notes=["DOCX is not a valid archive."], error_code="parse_error")
    try:
        from docx import Document

        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    except Exception:
        return _result(
            "failed",
            notes=["DOCX extraction failed; file may be malformed or unsupported."],
            indicators=indicators,
            error_code="parse_error",
        )
    return _result(
        "ok" if text.strip() else "no_text",
        text,
        [] if text.strip() else ["DOCX contains no extractable text."],
        indicators,
    )


def _extract_image(raw: bytes) -> dict:
    try:
        from PIL import Image, ImageOps
    except ImportError:
        return _result(
            "failed",
            notes=["Pillow is not installed; images cannot be inspected."],
            error_code="parser_unavailable",
        )
    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    try:
        with Image.open(io.BytesIO(raw)) as im:
            im.verify()
        with Image.open(io.BytesIO(raw)) as im:
            im = ImageOps.exif_transpose(im).convert("RGB")
    except Image.DecompressionBombError:
        return _result(
            "limit",
            notes=["Image dimensions exceed the safe decoding limit."],
            error_code="image_too_large",
        )
    except Exception:
        return _result(
            "failed",
            notes=["Image validation failed; file may be malformed or too large to decode safely."],
            error_code="parse_error",
        )
    try:
        import pytesseract
    except ImportError:
        return _result(
            "no_text",
            notes=[
                "Image validated, but OCR is unavailable (pytesseract/Tesseract not installed)."
            ],
            error_code="ocr_unavailable",
        )
    try:
        text = pytesseract.image_to_string(im, timeout=OCR_TIMEOUT_SECONDS)
    except Exception:
        return _result(
            "failed",
            notes=["OCR failed or timed out; image may be unclear or Tesseract unavailable."],
            error_code="ocr_failed",
        )
    return _result(
        "ok" if text.strip() else "no_text",
        text,
        [] if text.strip() else ["OCR found no readable text."],
    )


def extract(ext: str, raw: bytes) -> dict:
    if ext in {".txt", ".msg"}:
        text, notes = decode_text(raw)
        return _result("ok" if text.strip() else "no_text", text, notes)
    if ext == ".eml":
        return _extract_eml(raw)
    if ext == ".pdf":
        return _extract_pdf(raw)
    if ext == ".docx":
        return _extract_docx(raw)
    if ext in IMAGE_EXTS:
        return _extract_image(raw)
    return _result("failed", notes=["Unsupported file type."], error_code="unsupported_type")


def main() -> int:
    apply_limits()
    ext = sys.argv[1].lower() if len(sys.argv) > 1 else ""
    raw = sys.stdin.buffer.read()
    try:
        out = extract(ext, raw)
    except MemoryError:
        out = _result("limit", notes=["Parser exceeded its memory limit."], error_code="memory")
    except RecursionError:
        out = _result("failed", notes=["Parser recursion limit hit."], error_code="parse_error")
    except Exception:
        out = _result("failed", notes=["Unexpected parser failure."], error_code="parse_error")
    sys.stdout.write(json.dumps(out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
